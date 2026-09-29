"""Offline HTML and native-chart Excel exports, generated from one analysis."""
import io
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from analysis import COST_HEADERS, PROD_HEADERS, COST_SHEET, PROD_SHEET, DRIVERS, public_report

ROOT = Path(__file__).parent
PALETTE = ['2864A0','DA8742','33937D','855FA8','607384','C14D4D']


def safe_sheet(ws):
    """Force source text to stay text, including Excel formula-leading strings."""
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str):
                cell.data_type = 's'


def style(ws, header=1):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f'C{header+1}'
    for row in ws:
        for c in row:
            c.font = Font(name='Arial',size=10,color='172B42')
            c.alignment = Alignment(vertical='center')
            if isinstance(c.value,(float,int)):
                c.number_format = '#,##0.00;[Red](#,##0.00);"-"'
            if isinstance(c.value,datetime):
                c.number_format = 'yyyy-mm-dd'
        ws.row_dimensions[row[0].row].height = 23
    for c in ws[header]:
        c.fill = PatternFill('solid',fgColor='243D57')
        c.font = Font(name='Arial',size=10,bold=True,color='FFFFFF')
        c.alignment = Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[header].height = 34
    for col in ws.columns:
        letter=col[0].column_letter
        ws.column_dimensions[letter].width = min(32,max(16,max(len(str(c.value or '')) for c in list(col)[:60])*1.5))


def to_bytes(wb):
    stream=io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def template():
    wb=Workbook()
    ws=wb.active
    ws.title=COST_SHEET
    ws.append(COST_HEADERS)
    production=wb.create_sheet(PROD_SHEET)
    production.append(PROD_HEADERS)
    help_=wb.create_sheet('填写说明')
    for row in [
        ['项目','填写规则'],
        ['成本行','每行一条制造成本/期间费用记录；可同月同产品同科目多条。编号必须唯一。'],
        ['月份','Excel日期或YYYY-MM，如2025-01；数据首行为列名，不要改名。'],
        ['费用口径','制造成本或期间费用。期间费用只汇总，不计入产品单位成本。'],
        ['成本性态','固定、变动、混合。变动项按价格/投产单耗/良率拆分。'],
        ['金额','数量×单价=不含税金额（元），允许0.01元舍入差。输入纯数值，不用公式。'],
        ['产量行','每产品每月仅一行；投产=报废+合格入库；投产和合格量>0。'],
        ['边界','不处理在制品变动、负数冲回或停产分摊；此类情况需另行调整口径。'],
        ['计量','产品编码对应唯一名称和产量单位。同产品同科目须使用一致耗用单位/性态/类别。'],
        ['可选列','成本对象键、发生日期、成本中心、内部科目编码、归集方式、业务说明及工时不参与核心计算。'],
        ['缺月','缺少相邻上月时不计算环比，不把缺月当作0。'],
        ['多笔归集','同产品同月同科目按耗用量加权单价，不能简单平均价格。'],
        ['操作','填完两个原始表，保存为.xlsx后上传；下载示例可查看完整1000条模拟数据。'],
    ]:help_.append(row)
    for s in wb:style(s)
    help_.column_dimensions['B'].width=105
    return to_bytes(wb)


def choose(report, product=None, month=None):
    product = product or report['products'][0]['code']
    periods = [r['month'] for r in report['monthly'] if r['product']==product]
    if not periods:
        raise ValueError('产品不存在')
    month = month or max(periods)
    if month not in periods:
        raise ValueError('所选产品在该月份没有数据')
    return product,month


def calendar_months(months):
    first,last=min(months),max(months)
    y0,m0=map(int,first.split('-'));y1,m1=map(int,last.split('-'))
    return [f'{i//12:04d}-{i%12+1:02d}' for i in range(y0*12+m0-1,y1*12+m1)]


def excel(report, product=None, month=None):
    product,month=choose(report,product,month)
    wb=Workbook()
    cover=wb.active
    cover.title='成本图表'
    monthly=wb.create_sheet('月度成本分析')
    factors=wb.create_sheet('变动因素明细')
    data=wb.create_sheet('图表数据')
    source=report['_source']
    for sheet,h,rows in [(COST_SHEET,source['cost_headers'],source['cost_rows']),(PROD_SHEET,source['prod_headers'],source['prod_rows'])]:
        ws=wb.create_sheet(sheet);ws.append(h)
        for r in rows:ws.append([r.get(k) for k in h])
        safe_sheet(ws)
        tab=Table(displayName='CostInput' if sheet==COST_SHEET else 'ProductionInput',ref=ws.dimensions)
        tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(tab)
    categories=sorted({c for r in report['monthly'] for c in r['categories']})
    names={p['code']:p['name'] for p in report['products']}
    units={p['code']:p['unit'] for p in report['products']}
    monthly.append(['月份','产品编码','产品名称','产量单位','合格产量','投产数量','制造成本（元）','单位成本（元/产量单位）','固定费用/单位']+[f'{c}（元/产量单位）' for c in categories])
    for r in report['monthly']:monthly.append([r['month'],r['product'],names[r['product']],units[r['product']],r['good'],r['input'],r['total'],r['unit_cost'],r['fixed_unit']]+[r['categories'].get(c,0) for c in categories])
    factors.append(['月份','产品编码','科目','类别','性态','上月单位成本','本月单位成本','单位成本变动']+DRIVERS+['校验差额'])
    for comp in report['comparisons']:
        for s in comp['subjects']:factors.append([comp['month'],comp['product'],s['subject'],s['category'],s['nature'],s['previous_cost'],s['current_cost'],s['delta']]+[s['drivers'][d] for d in DRIVERS]+[s['delta']-sum(s['drivers'].values())])
    current=next(r for r in report['monthly'] if r['product']==product and r['month']==month)
    comp=next((r for r in report['comparisons'] if r['product']==product and r['month']==month),None)
    unit=units[product]
    cover.append(['成本趋势与原因分析'])
    cover.append([f'{names[product]} · {month} · 元/{unit} · 来源：{report["filename"]}'])
    cover.append(['当前单位成本',current['unit_cost'],'制造成本（元）',current['total'],'合格产量',current['good']])
    cover.append(['环比变动',comp['delta'] if comp else '无相邻上月','环比变动率',comp['rate'] if comp and comp['rate'] is not None else '不适用'])
    cover.cell(4,4).number_format='0.0%'
    cover.append(['说明：本文件是分析快照。修改原始数据后请重新上传生成；图表可在Excel中编辑。'])
    for i,s in enumerate(comp['summary'] if comp else ['无相邻上月，不计算环比原因。'],6):cover.cell(i,1,s)
    # Wide source area is intentionally separate from charts.
    same_units=[p for p in report['products'] if p['unit']==unit]
    data.append(['月份']+[p['name'] for p in same_units])
    periods=calendar_months(report['months'])
    for m in periods:
        data.append([m]+[next((r['unit_cost'] for r in report['monthly'] if r['product']==p['code'] and r['month']==m),None) for p in same_units])
    n=len(periods)+1
    selected_lookup={r['month']:r for r in report['monthly'] if r['product']==product}
    selected_periods=calendar_months(selected_lookup)
    selected=[selected_lookup.get(m) for m in selected_periods]
    start=n+4
    for c,v in enumerate(['月份']+categories+['合格产量','固定费用/单位'],1):data.cell(start,c,v)
    for i,(m,r) in enumerate(zip(selected_periods,selected),start+1):
        values=[m]+([r['categories'].get(c,0) for c in categories]+[r['good'],r['fixed_unit']] if r else [None]*(len(categories)+2))
        for j,v in enumerate(values,1):data.cell(i,j,v)
    driver_start=start+len(selected)+4
    data.cell(driver_start,1,'因素');data.cell(driver_start,2,f'影响（元/{unit}）')
    for i,d in enumerate(DRIVERS,driver_start+1):data.cell(i,1,d);data.cell(i,2,comp['drivers'][d] if comp else None)
    cat_start=driver_start+10
    data.cell(cat_start,1,'成本类别');data.cell(cat_start,2,f'环比贡献（元/{unit}）')
    for i,c in enumerate(categories,cat_start+1):data.cell(i,1,c);data.cell(i,2,comp['categories'].get(c,0) if comp else None)
    def add(kind,title,start,end,c1,c2,anchor,ytitle):
        chart=LineChart() if kind=='line' else BarChart()
        chart.title=title;chart.style=13;chart.y_axis.title=ytitle;chart.x_axis.title='月份' if kind=='line' else '因素'
        chart.add_data(Reference(data,min_col=c1,max_col=c2,min_row=start,max_row=end),titles_from_data=True)
        chart.set_categories(Reference(data,min_col=1,min_row=start+1,max_row=end))
        chart.width=22;chart.height=10;chart.legend.position='b';chart.display_blanks='gap'
        if kind=='bar':chart.type='col'
        for i,series in enumerate(chart.series):
            series.graphicalProperties.line.solidFill=PALETTE[i%len(PALETTE)]
            series.graphicalProperties.solidFill=PALETTE[i%len(PALETTE)]
        cover.add_chart(chart,anchor)
    add('line','同计量单位产品成本趋势',1,n,2,1+len(same_units),'A13',f'元/{unit}')
    add('line','所选产品成本构成趋势',start,start+len(selected),2,1+len(categories),'M13',f'元/{unit}')
    add('bar','环比核心因素贡献' if comp else '无相邻上月：无因素贡献',driver_start,driver_start+6,2,2,'A34',f'元/{unit}')
    add('bar','成本类别环比贡献' if comp else '无相邻上月：无类别贡献',cat_start,cat_start+len(categories),2,2,'M34',f'元/{unit}')
    add('line','合格产量趋势',start,start+len(selected),len(categories)+2,len(categories)+2,'A55',unit)
    add('line','固定费用分摊趋势',start,start+len(selected),len(categories)+3,len(categories)+3,'M55',f'元/{unit}')
    help_=wb.create_sheet('分析口径与提示')
    for row in [['项目','说明'],['期间费用（元）',report['period_expense_total']],['制造单位成本','制造成本÷合格产量，期间费用不计入，不进行跨产品平均。'],['分解顺序','变动项：价格→投入单耗→良率；固定/混合项：费用额→产量分摊。'],['交互项','分解顺序决定交互项归属，贡献合计等于单位成本变化。'],['加权价格','同月同产品同科目按耗用量加权单价；金额舍入差单列。'],['缺项/零耗用','变动科目缺项或耗用为零时按费用额和产量拆分，不推断价格。'],['局限','这是算术归因，真实经营原因需结合采购、工艺、维修、报废等证据。'],['静态快照','本工作簿不在修改原始表后自动重新执行Python分析，请重新上传。']]:help_.append(row)
    for w in report['warnings']:help_.append(['数据提示',w])
    for ws in wb:
        safe_sheet(ws);style(ws)
    cover.freeze_panes=None
    for i in range(1,25):cover.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width=11
    cover['A1'].font=Font(name='Arial',size=16,bold=True,color='243D57');cover['A1'].fill=PatternFill(fill_type=None)
    cover['D4'].number_format='0.0%'
    help_.column_dimensions['B'].width=110
    wb.active=0
    return to_bytes(wb)


def html_report(report,product=None,month=None):
    product,month=choose(report,product,month)
    page=(ROOT/'static/index.html').read_text()
    css=(ROOT/'static/style.css').read_text()
    js=(ROOT/'static/app.js').read_text()
    payload=json.dumps(public_report(report),ensure_ascii=False,allow_nan=False).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
    selection=json.dumps({'product':product,'month':month},ensure_ascii=False).replace('<','\\u003c')
    page=re.sub(r'<link\s+rel="stylesheet"\s+href="/static/style.css"\s*/?>',lambda _:f'<style>{css}</style>',page)
    page=re.sub(r'<script\s+src="/static/app.js"(?:\s+defer)?\s*></script>',lambda _:f'<script>window.COST_REPORT={payload};window.COST_SELECTION={selection};</script><script>{js}</script>',page)
    if '/static/app.js' in page or '/static/style.css' in page:
        raise RuntimeError('离线模板资源未完成内嵌')
    return page.encode('utf-8')
