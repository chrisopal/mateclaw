"""Validated manufacturing cost calculations. All monetary outputs use yuan."""
from __future__ import annotations

import io
import math
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime

from openpyxl import load_workbook

COST_SHEET = '成本原始明细'
PROD_SHEET = '产量原始数据'
COST_HEADERS = ['明细编号', '发生日期', '所属月份', '成本对象键', '产品编码', '产品名称', '成本中心', '内部科目编码', '成本科目', '成本类别', '费用口径', '成本性态', '耗用单位', '耗用数量', '不含税单价（元）', '不含税金额（元）', '归集方式', '业务说明']
PROD_HEADERS = ['成本对象键', '所属月份', '产品编码', '产品名称', '车间', '产量单位', '投产数量', '报废数量', '合格入库产量', '实际人工工时', '实际机器工时']
COST_REQUIRED = ['明细编号','所属月份','产品编码','产品名称','成本科目','成本类别','费用口径','成本性态','耗用单位','耗用数量','不含税单价（元）','不含税金额（元）']
PROD_REQUIRED = ['所属月份','产品编码','产品名称','产量单位','投产数量','报废数量','合格入库产量']
DRIVERS = ['价格变化','投入单耗变化','良率变化','费用额变化','产量摊薄','舍入差异']
MAX_BYTES = 10 * 1024 * 1024


class ValidationError(ValueError):
    def __init__(self, errors):
        self.errors = errors if isinstance(errors, list) else [errors]
        super().__init__('；'.join(self.errors))


def month(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m')
    match = re.fullmatch(r'(\d{4})[-/](\d{1,2})(?:[-/](\d{1,2}))?', str(value).strip())
    if not match:
        raise ValueError('月份须为Excel日期或 YYYY-MM')
    year, mon = int(match[1]), int(match[2])
    day = int(match[3] or 1)
    try:
        date(year, mon, day)
    except ValueError as exc:
        raise ValueError('月份或日期无效，请使用真实日期') from exc
    return f'{year:04d}-{mon:02d}'


def prior_month(value):
    y, m = map(int, value.split('-'))
    return f'{y-1:04d}-12' if m == 1 else f'{y:04d}-{m-1:02d}'


def number(value, field):
    if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value) or value < 0:
        raise ValueError(f'{field}须为非负数值，不可空白')
    return float(value)


def text(value, field):
    if value is None or not str(value).strip():
        raise ValueError(f'{field}不可空白')
    return str(value).strip()


def _table(wb, name, required):
    if name not in wb.sheetnames:
        raise ValidationError(f'缺少工作表：{name}')
    ws = wb[name]
    if (ws.max_row or 0) > 20001 or (ws.max_column or 0) > 100:
        raise ValidationError(f'{name}超过限制：最多20000行记录、100列')
    rows = ws.iter_rows(values_only=True)
    headers = [str(v).strip() if v is not None else '' for v in next(rows, [])]
    duplicate = {h for h in headers if h and headers.count(h) > 1}
    if duplicate:
        raise ValidationError(f'{name}有重复列名：{", ".join(sorted(duplicate))}')
    missing = set(required) - set(headers)
    if missing:
        raise ValidationError(f'{name}缺少列：{", ".join(sorted(missing))}')
    data = []
    for line, values in enumerate(rows, 2):
        if line > 20001 or len(values)>100:
            raise ValidationError(f'{name}超过限制：最多20000行记录、100列')
        if all(v is None for v in values):
            continue
        row = dict(zip(headers, values))
        if any(isinstance(row.get(k), str) and row[k].startswith('=') for k in required):
            raise ValidationError(f'{name}第{line}行包含公式，请粘贴为数值后上传')
        data.append((line, row))
    if not data:
        raise ValidationError(f'{name}没有数据；请填写模板或使用示例')
    return headers, data


def read_input(content):
    if len(content) > MAX_BYTES:
        raise ValidationError('文件超过10MB上限')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValidationError('Excel解压大小超过100MB或内部文件过多')
            if any('vbaproject' in x.filename.lower() for x in archive.infolist()):
                raise ValidationError('不支持带宏的工作簿，请另存为xlsx')
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            cost_headers, cost = _table(wb, COST_SHEET, COST_REQUIRED)
            prod_headers, production = _table(wb, PROD_SHEET, PROD_REQUIRED)
        finally:
            wb.close()
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError('无法读取Excel，请上传未加密的有效 .xlsx 文件') from exc
    return cost_headers, cost, prod_headers, production


def analyze(content, filename='成本数据.xlsx'):
    ch, cost_rows, ph, production_rows = read_input(content)
    errors, warnings, production, products = [], [], {}, {}
    for line, row in production_rows:
        try:
            p, m = text(row['产品编码'], '产品编码'), month(row['所属月份'])
            name, unit = text(row['产品名称'], '产品名称'), text(row['产量单位'], '产量单位')
            values = {k: number(row[k], k) for k in ['投产数量','报废数量','合格入库产量']}
            good, inp = values['合格入库产量'], values['投产数量']
            if good <= 0 or inp <= 0:
                raise ValueError('投产数量和合格入库产量必须大于0，停产月份请单独分析')
            if abs(inp - good - values['报废数量']) > 0.000001:
                raise ValueError('投产数量必须等于合格入库产量＋报废数量（本模板无在制品变动）')
            if (p, m) in production:
                raise ValueError(f'产品{p}月份{m}产量重复；请先合并为一行')
            if p in products and products[p] != {'code':p,'name':name,'unit':unit}:
                raise ValueError(f'产品{p}名称或产量单位跨月不一致')
            products[p] = {'code':p,'name':name,'unit':unit}
            production[p, m] = {'good':good,'input':inp}
        except (ValueError, TypeError) as exc:
            errors.append(f'{PROD_SHEET}第{line}行：{exc}')
    ids, signatures = set(), {}
    expenses = 0.
    groups = defaultdict(lambda: defaultdict(lambda: {'q':0.,'qp':0.,'amount':0.}))
    for line, row in cost_rows:
        try:
            rid = text(row['明细编号'],'明细编号')
            if rid in ids:
                raise ValueError(f'明细编号{rid}重复')
            ids.add(rid)
            m = month(row['所属月份'])
            q, price, amount = [number(row[k], k) for k in ['耗用数量','不含税单价（元）','不含税金额（元）']]
            if abs(round(q * price, 2) - amount) > 0.011:
                raise ValueError('不含税金额与耗用数量×单价不符（允许0.01元舍入差）')
            scope = text(row['费用口径'], '费用口径')
            if scope not in ['制造成本','期间费用']:
                raise ValueError('费用口径仅允许制造成本、期间费用')
            if scope == '期间费用':
                expenses += amount
                continue
            p = text(row['产品编码'],'产品编码')
            if (p, m) not in production:
                raise ValueError(f'没有产品{p}月份{m}的有效产量记录')
            if text(row['产品名称'],'产品名称') != products[p]['name']:
                raise ValueError('产品名称与产量表不一致')
            subject, category, nature, unit = [text(row[k],k) for k in ['成本科目','成本类别','成本性态','耗用单位']]
            if nature not in ['固定','变动','混合']:
                raise ValueError('成本性态仅允许固定、变动、混合')
            sig = (category, nature, unit)
            if (p, subject) in signatures and signatures[p, subject] != sig:
                raise ValueError(f'{subject}的类别/性态/耗用单位不一致，需统一或拆成独立科目')
            signatures[p, subject] = sig
            item = groups[p, m][subject]
            item.update(subject=subject,category=category,nature=nature,unit=unit)
            item['q'] += q
            item['qp'] += q * price
            item['amount'] += amount
        except (ValueError, TypeError) as exc:
            errors.append(f'{COST_SHEET}第{line}行：{exc}')
    for p, m in production:
        if (p, m) not in groups:
            errors.append(f'产品{p}月份{m}有产量但没有制造成本，不能将缺失成本当作0')
    if errors:
        raise ValidationError(errors[:50] + ([f'另有{len(errors)-50}处错误'] if len(errors)>50 else []))
    monthly, lookup = [], {}
    for (p, m), subjects in sorted(groups.items()):
        good, inp = production[p,m]['good'], production[p,m]['input']
        cats = defaultdict(float)
        fixed = total = 0.
        for s in subjects.values():
            cats[s['category']] += s['amount']/good
            total += s['amount']
            if s['nature'] == '固定':
                fixed += s['amount']
        row = {'product':p,'month':m,'good':good,'input':inp,'total':total,'unit_cost':total/good,'fixed_unit':fixed/good,'categories':dict(cats)}
        monthly.append(row)
        lookup[p,m] = row
    comparisons = []
    for cur in monthly:
        p, m = cur['product'], cur['month']
        prev = lookup.get((p, prior_month(m)))
        if not prev:
            warnings.append(f'{p} {m}无相邻上月数据：保留成本趋势，不计算环比原因。')
            continue
        g0, g1, n0, n1 = prev['good'], cur['good'], prev['input'], cur['input']
        before, after = groups[p,prev['month']], groups[p,m]
        drivers, categories, details = dict.fromkeys(DRIVERS,0.), defaultdict(float), []
        for subject in sorted(set(before)|set(after)):
            a, b = before.get(subject), after.get(subject)
            spec = a or b
            amt0, amt1 = (a or {}).get('amount',0.), (b or {}).get('amount',0.)
            unit0, unit1 = amt0/g0, amt1/g1
            parts = dict.fromkeys(DRIVERS,0.)
            if a and b and spec['nature']=='变动' and a['q']>0 and b['q']>0:
                q0,q1 = a['q'],b['q']
                p0,p1 = a['qp']/q0,b['qp']/q1
                parts['价格变化'] = q0/g0*(p1-p0)
                parts['投入单耗变化'] = p1*(q1/n1-q0/n0)*n0/g0
                parts['良率变化'] = p1*q1/n1*(n1/g1-n0/g0)
                parts['舍入差异'] = (amt1-b['qp'])/g1-(amt0-a['qp'])/g0
            else:
                parts['费用额变化'] = (amt1-amt0)/g0
                parts['产量摊薄'] = amt1*(1/g1-1/g0)
                if spec['nature']=='变动':
                    warnings.append(f'{p} {m} {subject}上月/本月缺项或耗用为0，按费用额及产量变化分解，不推断价格。')
            delta = unit1-unit0
            for key,value in parts.items():
                drivers[key] += value
            categories[spec['category']] += delta
            details.append({'subject':subject,'category':spec['category'],'nature':spec['nature'],'previous_cost':unit0,'current_cost':unit1,'delta':delta,'drivers':parts})
        delta = cur['unit_cost']-prev['unit_cost']
        check = delta-sum(drivers.values())
        if abs(check)>1e-7:
            raise ValidationError(f'{p} {m}变动分解校验未通过')
        ranked = sorted([(k,v) for k,v in drivers.items() if k!='舍入差异' and abs(v)>0.005],key=lambda x:abs(x[1]),reverse=True)
        summary = [f'{m}较{prev["month"]}单位成本{"增加" if delta>=0 else "减少"}{abs(delta):.2f}元/{products[p]["unit"]}。']
        summary += [f'{k}贡献{v:+.2f}元/{products[p]["unit"]}，{"推高" if v>0 else "降低"}单位成本。' for k,v in ranked[:3]]
        top = max(details,key=lambda s:abs(s['delta']))
        summary.append(f'变动最大的科目：{top["subject"]}（{top["delta"]:+.2f}元/{products[p]["unit"]}）。')
        comparisons.append({'product':p,'month':m,'previous_month':prev['month'],'previous_cost':prev['unit_cost'],'current_cost':cur['unit_cost'],'delta':delta,'rate':delta/prev['unit_cost'] if prev['unit_cost'] else None,'volume_rate':g1/g0-1,'drivers':drivers,'categories':dict(categories),'subjects':sorted(details,key=lambda s:abs(s['delta']),reverse=True),'check':check,'summary':summary})
    return {'filename':filename,'record_count':len(cost_rows),'period_expense_total':expenses,'warnings':list(dict.fromkeys(warnings)),'products':list(products.values()),'months':sorted({m for _,m in production}),'monthly':monthly,'comparisons':comparisons,'_source':{'cost_headers':ch,'cost_rows':[r for _,r in cost_rows],'prod_headers':ph,'prod_rows':[r for _,r in production_rows]}}


def public_report(report):
    return {k:v for k,v in report.items() if not k.startswith('_')}
