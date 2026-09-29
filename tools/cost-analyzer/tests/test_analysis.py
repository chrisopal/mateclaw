import io
import unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook

from analysis import analyze,ValidationError,COST_HEADERS,PROD_HEADERS,COST_SHEET,PROD_SHEET,public_report
from reports import excel,template,html_report


def fixture(cost=None,production=None):
    costs=cost if cost is not None else [
        ['A','2025-01','P1','产品A','材料','直接材料','制造成本','变动','kg',200,5,1000],
        ['B','2025-02','P1','产品A','材料','直接材料','制造成本','变动','kg',240,6,1440],
        ['C','2025-01','P1','产品A','折旧','制造费用','制造成本','固定','分摊系数',1,500,500],
        ['D','2025-02','P1','产品A','折旧','制造费用','制造成本','固定','分摊系数',1,500,500],
    ]
    prods=production if production is not None else [
        ['2025-01','P1','产品A','件',100,0,100],['2025-02','P1','产品A','件',120,20,100],
    ]
    wb=Workbook();ws=wb.active;ws.title=COST_SHEET
    ws.append(['明细编号','所属月份','产品编码','产品名称','成本科目','成本类别','费用口径','成本性态','耗用单位','耗用数量','不含税单价（元）','不含税金额（元）'])
    for r in costs:ws.append(r)
    ws=wb.create_sheet(PROD_SHEET);ws.append(['所属月份','产品编码','产品名称','产量单位','投产数量','报废数量','合格入库产量'])
    for r in prods:ws.append(r)
    stream=io.BytesIO();wb.save(stream)
    return stream.getvalue()


def mutate(data,fn):
    wb=load_workbook(io.BytesIO(data));fn(wb)
    stream=io.BytesIO();wb.save(stream);return stream.getvalue()


class AnalysisTests(unittest.TestCase):
    def test_price_and_yield_bridge(self):
        r=analyze(fixture());c=r['comparisons'][0]
        self.assertAlmostEqual(c['previous_cost'],15)
        self.assertAlmostEqual(c['current_cost'],19.4)
        self.assertAlmostEqual(c['drivers']['价格变化'],2)
        self.assertAlmostEqual(c['drivers']['投入单耗变化'],0)
        self.assertAlmostEqual(c['drivers']['良率变化'],2.4)
        self.assertAlmostEqual(c['check'],0)

    def test_multiple_rows_use_weighted_price(self):
        data=mutate(fixture(),lambda w:w[COST_SHEET].append(['E','2025-02','P1','产品A','材料','直接材料','制造成本','变动','kg',60,10,600]))
        c=analyze(data)['comparisons'][0]
        self.assertAlmostEqual(c['drivers']['价格变化'],3.6)
        self.assertAlmostEqual(c['drivers']['投入单耗变化'],3.4)
        self.assertAlmostEqual(c['drivers']['良率变化'],3.4)
        self.assertAlmostEqual(c['delta'],10.4)

    def test_volume_allocation(self):
        data=mutate(fixture(),lambda w:[setattr(w[PROD_SHEET][cell],'value',value) for cell,value in [('E3',80),('F3',0),('G3',80)]])
        c=analyze(data)['comparisons'][0]
        self.assertAlmostEqual(c['drivers']['产量摊薄'],1.25)
        self.assertAlmostEqual(c['check'],0)

    def test_expenses_excluded(self):
        data=mutate(fixture(),lambda w:w[COST_SHEET].append(['E','2025-02','总部','公共','管理','管理费用','期间费用','固定','月',1,5000,5000]))
        r=analyze(data)
        self.assertEqual(r['period_expense_total'],5000)
        self.assertAlmostEqual(r['comparisons'][0]['current_cost'],19.4)

    def test_missing_month_no_comparison(self):
        def edit(w):
            for r in [3,5]:w[COST_SHEET].cell(r,2,'2025-03')
            w[PROD_SHEET]['A3']='2025-03'
        r=analyze(mutate(fixture(),edit))
        self.assertEqual(r['comparisons'],[])
        self.assertTrue(any('相邻上月' in w for w in r['warnings']))
        exported=load_workbook(io.BytesIO(excel(r)))
        self.assertEqual(exported['图表数据']['A3'].value,'2025-02')
        self.assertIsNone(exported['图表数据']['B3'].value)

    def test_new_subject_scope_change(self):
        data=mutate(fixture(),lambda w:w[COST_SHEET].append(['E','2025-02','P1','产品A','辅料','直接材料','制造成本','变动','kg',10,2,20]))
        c=analyze(data)['comparisons'][0]
        self.assertAlmostEqual(c['drivers']['费用额变化'],.2)
        self.assertAlmostEqual(c['check'],0)

    def test_zero_previous_cost_rate_unavailable(self):
        def edit(w):
            for r in [2,4]:w[COST_SHEET].cell(r,11,0);w[COST_SHEET].cell(r,12,0)
        c=analyze(mutate(fixture(),edit))['comparisons'][0]
        self.assertIsNone(c['rate'])

    def test_bad_inputs(self):
        cases=[
            (lambda w:setattr(w[COST_SHEET]['A1'],'value','错误列名'),'缺少列'),
            (lambda w:w[PROD_SHEET].append(['2025-01','P1','产品A','件',100,0,100]),'产量重复'),
            (lambda w:setattr(w[COST_SHEET]['L2'],'value',999),'金额与'),
            (lambda w:setattr(w[PROD_SHEET]['G2'],'value',0),'必须大于0'),
            (lambda w:setattr(w[COST_SHEET]['J2'],'value','=100+100'),'包含公式'),
            (lambda w:setattr(w[COST_SHEET]['A3'],'value','A'),'编号A重复'),
            (lambda w:setattr(w[COST_SHEET]['I3'],'value','吨'),'耗用单位不一致'),
            (lambda w:setattr(w[COST_SHEET]['K2'],'value',-5),'非负数值'),
            (lambda w:setattr(w[PROD_SHEET]['A3'],'value','2025-13'),'月份或日期无效'),
            (lambda w:w[COST_SHEET].delete_rows(3),'没有制造成本'),
        ]
        # Last case must remove both costs for the second month.
        cases[-1]=(lambda w:[w[COST_SHEET].delete_rows(r) for r in [5,3]],'没有制造成本')
        for fn,msg in cases:
            with self.subTest(msg=msg),self.assertRaises(ValidationError) as raised:analyze(mutate(fixture(),fn))
            self.assertIn(msg,str(raised.exception))

    def test_invalid_workbook_and_empty_template(self):
        with self.assertRaises(ValidationError):analyze(b'not excel')
        with self.assertRaisesRegex(ValidationError,'没有数据'):analyze(template())

    def test_export_roundtrip(self):
        r=analyze(fixture());data=excel(r,'P1','2025-02')
        wb=load_workbook(io.BytesIO(data))
        self.assertEqual(len(wb['成本图表']._charts),6)
        self.assertAlmostEqual(wb['月度成本分析']['H3'].value,19.4)
        self.assertAlmostEqual(wb['成本图表']['B4'].value,4.4)
        r2=analyze(data)
        self.assertEqual(r2['monthly'],r['monthly'])
        self.assertNotIn('_source',public_report(r))

    def test_text_never_becomes_formula_in_export(self):
        r=analyze(fixture());r['filename']='=HYPERLINK("http://example.com")'
        r['_source']['cost_rows'][0]['业务说明']='=1+1'
        r['_source']['cost_headers'].append('业务说明')
        w=load_workbook(io.BytesIO(excel(r)))
        self.assertEqual(w[COST_SHEET].cell(2,13).data_type,'s')

    def test_offline_resources_and_escaped_data(self):
        r=analyze(fixture(),'</script><img src=x onerror=alert(1)>')
        page=html_report(r,'P1','2025-02').decode()
        self.assertNotIn('src="/static/',page)
        self.assertNotIn('href="/static/',page)
        self.assertIn('window.COST_REPORT=',page)
        self.assertIn('\\u003c/script',page)

    def test_original_1000_rows(self):
        p=Path(__file__).parents[1]/'examples/制造企业成本示例.xlsx'
        r=analyze(p.read_bytes())
        self.assertEqual((r['record_count'],len(r['monthly']),len(r['comparisons'])),(1000,48,44))
        self.assertAlmostEqual(sum(x['total'] for x in r['monthly']),10684563.08)
        c=next(c for c in r['comparisons'] if c['product']=='P003' and c['month']=='2025-02')
        self.assertAlmostEqual(c['delta'],19.234743177422786)
        self.assertAlmostEqual(c['drivers']['产量摊薄'],19.056657832533485)
        self.assertTrue(all(abs(c['check'])<1e-8 for c in r['comparisons']))


if __name__=='__main__':unittest.main()
